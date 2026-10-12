"""Read-only settlement inspection and independent spreadsheet evidence."""

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from simulator.t108_contract import RunManifest, expected_segments, save_json
from src.modules.billing.models import Invoice, InvoiceLine
from src.modules.charging.models import ChargingSession
from src.modules.wallet.models import Wallet, WalletLedger


def compare(manifest: RunManifest, participant, trace: dict, actual: dict) -> dict:
    """Reject missing/extra sessions, lines, debits, stale or inconsistent balances."""
    samples = [
        (datetime.fromisoformat(r["at"]), Decimal(r["wh"])) for r in trace["samples"]
    ]
    expected = expected_segments(manifest, samples)
    total = sum(row["amount_vnd"] for row in expected)
    failures = []
    if len(actual["sessions"]) != 1 or actual.get("session_id") != trace.get(
        "session_id"
    ):
        failures.append("session_count_or_identity")
    if actual.get("driver_id") != str(participant.driver_id) or actual.get(
        "station_id"
    ) != str(manifest.station_id):
        failures.append("invoice_ownership")
    if (
        actual.get("started_at") != manifest.started_at.isoformat()
        or actual.get("ended_at") != manifest.ended_at.isoformat()
    ):
        failures.append("session_timestamps")
    if (
        actual.get("energy_wh") is None
        or Decimal(actual["energy_wh"]) != manifest.energy_wh
    ):
        failures.append("session_energy")
    if not actual.get("settled") or actual.get("settlement_reason") != "debited":
        failures.append("settlement_missing")
    if actual.get("invoice_count") != 1 or actual.get("total_vnd") != total:
        failures.append("invoice_total_or_count")
    if actual.get("lines") != [
        row
        | {
            "tariff_id": str(manifest.tariff_id),
            "line_type": "energy",
            "band_label": f"T108-{i:02d}",
        }
        for i, row in enumerate(expected, 1)
    ]:
        # Decimal text scale is normalized by the collector.
        failures.append("invoice_segments")
    if actual.get("debit_amounts") != [-total] or actual.get("debit_wallets") != [
        str(participant.wallet_id)
    ]:
        failures.append("debit_not_exactly_once")
    if actual.get("balance_vnd") != participant.balance_before_vnd - total:
        failures.append("balance_difference")
    if actual.get("balance_vnd") != actual.get("ledger_sum_vnd"):
        failures.append("ledger_cache_mismatch")
    if actual.get("topup_amounts") != [manifest.seed_amount_vnd]:
        failures.append("seed_receipt")
    if manifest.mode == "staging" and abs(trace.get("start_delay_seconds", 999)) > 2:
        failures.append("start_did_not_match_three_minute_schedule")
    return {
        "index": participant.index,
        "code": participant.charge_point_code,
        "session_id": trace.get("session_id"),
        "expected_segments": expected,
        "expected_total_vnd": total,
        "balance_before_vnd": participant.balance_before_vnd,
        "actual": actual,
        "failures": failures,
        "passed": not failures,
    }


def decimal_text(value) -> str:
    return format(Decimal(value).normalize(), "f") if value is not None else ""


async def inspect_participant(
    db: AsyncSession, manifest: RunManifest, participant
) -> dict:
    sessions = list(
        await db.scalars(
            select(ChargingSession)
            .where(ChargingSession.charge_point_id == participant.charge_point_id)
            .order_by(ChargingSession.id)
        )
    )
    result = {"sessions": [s.id for s in sessions]}
    if len(sessions) != 1:
        return result
    charging = sessions[0]
    invoices = list(
        await db.scalars(select(Invoice).where(Invoice.session_id == charging.id))
    )
    wallet = await db.get(Wallet, participant.wallet_id)
    debits = list(
        await db.scalars(
            select(WalletLedger).where(
                WalletLedger.entry_type == "charging_debit",
                WalletLedger.reference_id == str(charging.id),
            )
        )
    )
    topups = list(
        await db.scalars(
            select(WalletLedger).where(
                WalletLedger.entry_type == "manual_topup",
                WalletLedger.reference_id == participant.seed_receipt_code,
            )
        )
    )
    ledger_sum = await db.scalar(
        select(func.coalesce(func.sum(WalletLedger.amount_vnd), 0)).where(
            WalletLedger.wallet_id == participant.wallet_id
        )
    )
    result.update(
        session_id=charging.id,
        started_at=charging.started_at.isoformat(),
        ended_at=charging.ended_at.isoformat() if charging.ended_at else None,
        energy_wh=decimal_text(charging.energy_kwh * 1000)
        if charging.energy_kwh is not None
        else None,
        settled=charging.settlement_completed_at is not None,
        settlement_reason=charging.settlement_reason,
        invoice_count=len(invoices),
        balance_vnd=wallet.balance_vnd if wallet else None,
        ledger_sum_vnd=int(ledger_sum),
        debit_amounts=[d.amount_vnd for d in debits],
        debit_wallets=[str(d.wallet_id) for d in debits],
        topup_amounts=[t.amount_vnd for t in topups],
    )
    if invoices:
        invoice = invoices[0]
        lines = list(
            await db.scalars(
                select(InvoiceLine)
                .where(InvoiceLine.invoice_id == invoice.id)
                .order_by(InvoiceLine.started_at, InvoiceLine.line_type)
            )
        )
        result.update(
            total_vnd=invoice.total_vnd,
            driver_id=str(invoice.driver_id),
            station_id=str(invoice.station_id),
            lines=[
                {
                    "started_at": line.started_at.isoformat(),
                    "ended_at": line.ended_at.isoformat(),
                    "energy_wh": decimal_text(line.energy_wh),
                    "rate_vnd": line.rate_vnd,
                    "amount_vnd": line.amount_vnd,
                    "tariff_id": str(line.tariff_id),
                    "line_type": line.line_type,
                    "band_label": line.band_label,
                }
                for line in lines
            ],
        )
    return result


def write_evidence(
    directory: Path, manifest: RunManifest, traces: dict, rows: list[dict]
) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties

    report = {
        "run_id": str(manifest.run_id),
        "station_id": str(manifest.station_id),
        "source_revision": manifest.source_revision,
        "environment": manifest.mode,
        "count": len(rows),
        "passed": len(rows) == len(manifest.participants)
        and all(r["passed"] for r in rows),
        "staging_acceptance_candidate": manifest.mode == "staging"
        and len(rows) == 20
        and all(r["passed"] for r in rows),
        "rows": rows,
    }
    save_json(directory / "report.json", report)
    book = Workbook()
    summary = book.active
    summary.title = "Reconciliation"
    summary.append(
        [
            "Run",
            "Charger",
            "Session",
            "kWh",
            "Expected VND",
            "Invoice VND",
            "Before",
            "After",
            "Ledger sum",
            "Debit count",
            "Passed",
            "Failures",
        ]
    )
    sheet = book.create_sheet("Independent answers")
    sheet.append(
        [
            "Charger",
            "Segment",
            "Start",
            "End",
            "Meter start Wh",
            "Meter end Wh",
            "kWh formula",
            "Rate",
            "Money formula",
            "Expected cached",
            "Actual VND",
        ]
    )
    raw = book.create_sheet("Sent readings")
    raw.append(["Charger", "Timestamp", "Cumulative Wh"])
    for row in rows:
        trace = traces[str(row["index"])]
        sample_rows = {}
        for sample in trace["samples"]:
            raw.append([row["code"], sample["at"], float(Decimal(sample["wh"]))])
            sample_rows[sample["at"]] = raw.max_row
        for index, segment in enumerate(row["expected_segments"], 1):
            n = sheet.max_row + 1
            left, right = (
                sample_rows[segment["started_at"]],
                sample_rows[segment["ended_at"]],
            )
            actual_lines = row["actual"].get("lines", [])
            sheet.append(
                [
                    row["code"],
                    index,
                    segment["started_at"],
                    segment["ended_at"],
                    f"='Sent readings'!C{left}",
                    f"='Sent readings'!C{right}",
                    f"=(F{n}-E{n})/1000",
                    segment["rate_vnd"],
                    f"=ROUND(G{n}*H{n},0)",
                    segment["amount_vnd"],
                    actual_lines[index - 1]["amount_vnd"]
                    if len(actual_lines) >= index
                    else None,
                ]
            )
        actual = row["actual"]
        summary.append(
            [
                str(manifest.run_id),
                row["code"],
                row["session_id"],
                manifest.energy_wh / 1000,
                row["expected_total_vnd"],
                actual.get("total_vnd"),
                row["balance_before_vnd"],
                actual.get("balance_vnd"),
                actual.get("ledger_sum_vnd"),
                len(actual.get("debit_amounts", [])),
                row["passed"],
                ", ".join(row["failures"]),
            ]
        )
    for worksheet in book:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column in worksheet.columns:
            worksheet.column_dimensions[column[0].column_letter].width = min(
                60, max(14, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    book.calculation = CalcProperties(fullCalcOnLoad=True)
    book.save(directory / "reconciliation.xlsx")
    return report
